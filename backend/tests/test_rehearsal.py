"""Regression checks for failures found during the live demo rehearsal."""

from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from fastapi.testclient import TestClient
from app.services.document_service import extract_document, EmptyDocxError
from app.services import campaign_service as campaigns
from app.services.video_cache import cached_video_captions, cached_video_plan
from app.services.campaign_service import write_json


class RehearsalTests(unittest.TestCase):
    def setUp(self):
        self.storage = Path(__file__).resolve().parents[1] / "storage"
        self.temp = tempfile.TemporaryDirectory(prefix=".rehearsal-test-", dir=self.storage)
        self.folder = Path(self.temp.name)

    def tearDown(self):
        assert self.folder.resolve().is_relative_to(self.storage.resolve())
        self.temp.cleanup()

    def test_word_tables_preserve_order_and_merged_cells(self):
        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Workshop details"
        table.cell(1, 0).text = "Capacity"
        table.cell(1, 1).text = "24 participants"
        doc.add_paragraph("After table")
        path = self.folder / "table.docx"
        doc.save(path)
        result = extract_document(path, path.name, "docx")
        self.assertEqual(
            result["text"], "Before table\nWorkshop details\nCapacity\t24 participants\nAfter table"
        )
        table_only = Document()
        table_only.add_table(rows=1, cols=1).cell(0, 0).text = "Table-only text"
        table_only.save(path)
        self.assertEqual(extract_document(path, path.name, "docx")["text"], "Table-only text")

    def test_empty_word_upload_returns_helpful_422(self):
        from app.main import app

        source = BytesIO()
        Document().save(source)
        with (
            patch("app.api.upload.STORAGE_DIR", self.folder),
            patch("app.api.upload.find_run", return_value=None),
        ):
            with TestClient(app) as client:
                response = client.post(
                    "/api/upload",
                    files={
                        "file": (
                            "empty.docx",
                            source.getvalue(),
                            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        )
                    },
                )
        self.assertEqual(response.status_code, 422)
        self.assertIn("no readable text", response.json()["detail"])
        self.assertFalse(list(self.folder.iterdir()))

    def seed_campaign(self):
        write_json(
            self.folder,
            "extracted_text.json",
            {"file_type": "docx", "filename": "fixture.docx", "text": "Source"},
        )
        write_json(self.folder, "analysis.json", {"title": "Fixture"})
        write_json(
            self.folder,
            "master_content.json",
            {
                "headline": "Fixture",
                "lead": "Lead",
                "body": "Body",
                "press_release": "Release",
                "short_summary": "Summary",
            },
        )
        write_json(
            self.folder,
            "platforms.json",
            {"language": "en", "items": [], "platforms": [], "content_types": ["text"]},
        )

    def test_reopen_does_not_label_hungarian_master_as_old_english_posts(self):
        self.seed_campaign()
        write_json(self.folder, "output_language.json", {"language": "hu"})
        write_json(
            self.folder,
            "workspace.json",
            {"language": "hu", "content_types": ["text"], "platforms": ["website"]},
        )
        (self.folder / "video").mkdir()
        write_json(
            self.folder / "video",
            "status.json",
            {"language": "en", "stage": "ready", "url": "old-en-video"},
        )
        result = campaigns.bundle(self.folder)
        self.assertEqual(result["combined"]["output_language"], "hu")
        self.assertEqual(result["preview"]["language"], "en")
        self.assertIsNone(result["video_status"])

    def test_export_pairs_current_video_with_its_captions_and_plan(self):
        self.seed_campaign()
        video = self.folder / "video"
        video.mkdir()
        (video / "final_en.mp4").write_bytes(b"english-video-fixture")
        (video / "captions_en.srt").write_text("English captions")
        (video / "captions.srt").write_text("Hungarian captions")
        write_json(
            video, "video_plan_en.json", {"language": "en", "narration": "English narration"}
        )
        write_json(video, "video_plan.json", {"language": "hu", "narration": "Hungarian narration"})
        self.assertEqual(cached_video_plan(self.folder)["narration"], "English narration")
        self.assertEqual(cached_video_captions(self.folder).name, "captions_en.srt")
        payload, _ = campaigns.export_campaign(self.folder, True)
        with ZipFile(BytesIO(payload)) as archive:
            self.assertEqual(archive.read("media/video/captions_en.srt"), b"English captions")
            self.assertNotIn("media/video/captions.srt", archive.namelist())
        (video / "captions_en.srt").unlink()
        self.assertIsNone(cached_video_captions(self.folder))


if __name__ == "__main__":
    unittest.main()
